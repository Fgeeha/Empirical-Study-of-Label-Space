#!/usr/bin/env python3
"""Assemble paper/mdpi_ai/main.md into the MDPI Word template (ai-template.dot,
converted to submission/ai-template.docx with LibreOffice).

    python build_mdpi_docx.py            # -> submission/Kolesnikov_Kravets_AI.docx

Pipeline: main.md is split into front matter / body / back matter / references /
figure captions; body+back+refs go through pandoc (inline formatting, tables,
lists); the resulting paragraphs and tables are transplanted into the template
with MDPI paragraph styles, template three-line tables, figures inserted after
their first mention, and auto-numbered references. Needs pandoc and python-docx.
"""

import copy
from pathlib import Path
import re
import subprocess
import tempfile

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
MD = HERE / 'main.md'
TEMPLATE = HERE / 'submission/ai-template.docx'
OUT = HERE / 'submission/Kolesnikov_Kravets_AI.docx'
AUTHORS = [('Nikita S. Kolesnikov', '1,*'), ('Alla G. Kravets', '1')]
AFFILIATIONS = [
    '1\tDepartment of Computer-Aided Design and Search Engineering, Volgograd State '
    'Technical University, 28 Lenin Avenue, Volgograd 400005, Russia; '
    'kolesnikov.nikitavlg@gmail.com (N.S.K.); allagkravets@yandex.ru (A.G.K.)',
    '*\tCorrespondence: kolesnikov.nikitavlg@gmail.com',
]
W = qn


# ----------------------------------------------------------------------------- md
def parse_md(text: str) -> dict:
    lines = text.splitlines()
    idx = {ln.strip(): i for i, ln in enumerate(lines) if ln.startswith('## ')}
    title = next(ln[2:].strip() for ln in lines if ln.startswith('# '))
    i_abs, i_intro = idx['## Abstract'], idx['## 1. Introduction']
    abs_block = [
        ln for ln in lines[i_abs + 1 : i_intro] if ln.strip() and ln.strip() != '---'
    ]
    abstract = ' '.join(ln for ln in abs_block if not ln.startswith('**Keywords:**'))
    keywords = next(ln for ln in abs_block if ln.startswith('**Keywords:**'))
    keywords = keywords.replace('**Keywords:**', '').strip()
    i_back, i_refs, i_figs = (
        idx['## Back Matter'],
        idx['## References'],
        idx['## Figure Captions'],
    )
    body = [ln for ln in lines[i_intro:i_back] if ln.strip() != '---']
    back = [ln for ln in lines[i_back + 1 : i_refs] if ln.strip() != '---']
    refs = [ln for ln in lines[i_refs + 1 : i_figs] if ln.strip() != '---']
    figs = {}
    for ln in lines[i_figs + 1 :]:
        m = re.match(r'\*\*Figure (\d+)\.\*\*\s*(.*)', ln.strip())
        if m:
            cap = m.group(2)
            f = re.search(r'File: `([^`]+)`', cap)
            cap = re.sub(r'\s*File: `[^`]+`.*$', '', cap).strip()
            figs[int(m.group(1))] = (HERE / f.group(1), cap)
    md = (
        '\n'.join(body)
        + '\n\n## Back Matter\n\n'
        + '\n'.join(back)
        + '\n\n## References\n\n'
        + '\n'.join(refs)
        + '\n'
    )
    return dict(title=title, abstract=abstract, keywords=keywords, md=md, figs=figs)


# ---------------------------------------------------------------------- helpers
def set_runs(p: Paragraph, parts) -> None:
    """parts: list of (text, dict(bold=, italic=, sup=))."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for h in p._p.findall(W('w:hyperlink')):
        p._p.remove(h)
    for text, fmt in parts:
        r = p.add_run(text)
        r.bold = fmt.get('bold')
        r.italic = fmt.get('italic')
        if fmt.get('sup'):
            r.font.superscript = True


def md_inline(text: str):
    """Minimal **bold** / *italic* / `code` / ^sup^ parser for captions."""
    parts, pos = [], 0
    for m in re.finditer(r'\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\^(.+?)\^', text):
        if m.start() > pos:
            parts.append((text[pos : m.start()], {}))
        if m.group(1):
            parts.append((m.group(1), {'bold': True}))
        elif m.group(2):
            parts.append((m.group(2), {'italic': True}))
        elif m.group(3):
            parts.append((m.group(3), {}))
        else:
            parts.append((m.group(4), {'sup': True}))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], {}))
    return parts


def get_ppr(p_el):
    ppr = p_el.find(W('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p_el.insert(0, ppr)
    return ppr


class Builder:
    def __init__(self, dst: docx.document.Document):
        self.dst = dst
        self.body = dst.element.body
        self.sect = self.body.find(W('w:sectPr'))
        self.styles = {s.name: s.style_id for s in dst.styles}
        tbls = self.body.findall(W('w:tbl'))
        self.proto_tbl = tbls[1]  # template "Table 1" (three-line table)
        self.numbering = dst.part.numbering_part.element

    def style_id(self, name: str) -> str:
        return self.styles[name]

    def append(self, el) -> None:
        self.sect.addprevious(el)

    def new_par(self, style: str, parts=()):
        p = self.dst.add_paragraph(style=style)  # add_paragraph inserts before sectPr
        set_runs(p, parts)
        return p

    # numbering restart for a paragraph carrying a style with numPr
    def restart_num(self, p_el, style_name: str) -> None:
        st = self.dst.styles[style_name].element
        sty_num = st.pPr.find(W('w:numPr')).find(W('w:numId')).get(W('w:val'))
        base = next(
            n
            for n in self.numbering.findall(W('w:num'))
            if n.get(W('w:numId')) == sty_num
        )
        abs_id = base.find(W('w:abstractNumId')).get(W('w:val'))
        new_id = str(
            max(int(n.get(W('w:numId'))) for n in self.numbering.findall(W('w:num')))
            + 1
        )
        num = OxmlElement('w:num')
        num.set(W('w:numId'), new_id)
        a = OxmlElement('w:abstractNumId')
        a.set(W('w:val'), abs_id)
        num.append(a)
        ov = OxmlElement('w:lvlOverride')
        ov.set(W('w:ilvl'), '0')
        so = OxmlElement('w:startOverride')
        so.set(W('w:val'), '1')
        ov.append(so)
        num.append(ov)
        self.numbering.append(num)
        ppr = get_ppr(p_el)
        numpr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl')
        il.set(W('w:val'), '0')
        ni = OxmlElement('w:numId')
        ni.set(W('w:val'), new_id)
        numpr.append(il)
        numpr.append(ni)
        ppr.append(numpr)

    def clean_par(self, p_el, style_name: str) -> None:
        ppr = get_ppr(p_el)
        for tag in ('w:numPr', 'w:pStyle', 'w:ind', 'w:spacing', 'w:jc'):
            for e in ppr.findall(W(tag)):
                ppr.remove(e)
        ps = OxmlElement('w:pStyle')
        ps.set(W('w:val'), self.style_id(style_name))
        ppr.insert(0, ps)
        for h in p_el.findall(W('w:hyperlink')):
            for child in list(h):
                h.addprevious(child)
            p_el.remove(h)
        for r in p_el.iter(W('w:r')):
            rpr = r.find(W('w:rPr'))
            if rpr is not None:
                for e in rpr.findall(W('w:rStyle')):
                    rpr.remove(e)
        for bm in list(p_el.iter(W('w:bookmarkStart'))) + list(
            p_el.iter(W('w:bookmarkEnd'))
        ):
            bm.getparent().remove(bm)

    def add_table(self, src_tbl, src_doc) -> None:
        data = [
            [tc.findall(W('w:p')) for tc in tr.findall(W('w:tc'))]
            for tr in src_tbl.findall(W('w:tr'))
        ]
        ncols = max(len(r) for r in data)
        tbl = copy.deepcopy(self.proto_tbl)
        tblpr = tbl.find(W('w:tblPr'))
        for e in tblpr.findall(W('w:bidiVisual')):
            tblpr.remove(e)
        trs = tbl.findall(W('w:tr'))
        proto_head, proto_body = trs[0], trs[1]
        for tr in trs:
            tbl.remove(tr)
        grid = tbl.find(W('w:tblGrid'))
        total = sum(int(g.get(W('w:w'))) for g in grid.findall(W('w:gridCol')))
        for g in grid.findall(W('w:gridCol')):
            grid.remove(g)
        for _ in range(ncols):
            g = OxmlElement('w:gridCol')
            g.set(W('w:w'), str(total // ncols))
            grid.append(g)
        for ri, row in enumerate(data):
            tr = copy.deepcopy(proto_head if ri == 0 else proto_body)
            tcs = tr.findall(W('w:tc'))
            while len(tcs) < ncols:
                tr.append(copy.deepcopy(tcs[-1]))
                tcs = tr.findall(W('w:tc'))
            while len(tcs) > ncols:
                tr.remove(tcs[-1])
                tcs = tr.findall(W('w:tc'))
            for ci, tc in enumerate(tcs):
                tcw = tc.find(W('w:tcPr')).find(W('w:tcW'))
                if tcw is not None:
                    tcw.set(W('w:w'), str(total // ncols))
                    tcw.set(W('w:type'), 'dxa')
                ps = tc.findall(W('w:p'))
                for extra in ps[1:]:
                    tc.remove(extra)
                target = ps[0]
                for r in target.findall(W('w:r')):
                    target.remove(r)
                if ci < len(row):
                    for sp in row[ci]:
                        for r in sp.findall(W('w:r')):
                            nr = copy.deepcopy(r)
                            rpr = nr.find(W('w:rPr'))
                            if rpr is not None:
                                for e in rpr.findall(W('w:rStyle')):
                                    rpr.remove(e)
                            target.append(nr)
            tbl.append(tr)
        self.append(tbl)
        self.new_par('MDPI_3.2_text_no_indent', [('', {})])  # spacing after table

    def add_figure(self, n: int, path: Path, caption: str, width_cm: float) -> None:
        p = self.dst.add_paragraph(style='MDPI_5.2_figure')
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.new_par(
            'MDPI_5.1_figure_caption',
            [(f'Figure {n}. ', {'bold': True})] + md_inline(caption),
        )


# ------------------------------------------------------------------------- main
def main() -> None:
    info = parse_md(MD.read_text())
    with tempfile.TemporaryDirectory() as td:
        md_path, dx_path = Path(td) / 'body.md', Path(td) / 'body.docx'
        md_path.write_text(info['md'])
        subprocess.run(
            [
                'pandoc',
                str(md_path),
                '-f',
                'markdown',
                '-t',
                'docx',
                '--wrap=none',
                '-o',
                str(dx_path),
            ],
            check=True,
        )
        src = docx.Document(str(dx_path))
    dst = docx.Document(str(TEMPLATE))
    # The LibreOffice conversion of ai-template.dot marks the section, several styles
    # and the tables as right-to-left; strip those flags (the paper is LTR English).
    for sect in dst.sections:
        for tag in ('w:bidi', 'w:textDirection'):
            for e in sect._sectPr.findall(W(tag)):
                sect._sectPr.remove(e)
    for st in dst.styles.element.iter(W('w:style')):
        for pr_tag, flag in (('w:pPr', 'w:bidi'), ('w:rPr', 'w:rtl')):
            pr = st.find(W(pr_tag))
            if pr is not None:
                for e in pr.findall(W(flag)):
                    pr.remove(e)
    for tbl in dst.element.body.iter(W('w:tbl')):
        tp = tbl.find(W('w:tblPr'))
        for e in tp.findall(W('w:bidiVisual')):
            tp.remove(e)
    b = Builder(dst)
    src_style_names = {s.style_id: s.name for s in src.styles}

    # numbering formats of the pandoc document
    npart = src.part.numbering_part.element
    abs_fmt = {}
    for a in npart.findall(W('w:abstractNum')):
        lvl = a.find(W('w:lvl'))
        abs_fmt[a.get(W('w:abstractNumId'))] = lvl.find(W('w:numFmt')).get(W('w:val'))
    numfmt = {
        n.get(W('w:numId')): abs_fmt[n.find(W('w:abstractNumId')).get(W('w:val'))]
        for n in npart.findall(W('w:num'))
    }

    # ---- front matter in place
    body_els = list(b.body)
    pars = [Paragraph(e, dst) for e in body_els if e.tag == W('w:p')]
    set_runs(pars[0], [('Article', {})])
    set_runs(pars[1], [(info['title'], {})])
    auth_parts = []
    for i, (name, sup) in enumerate(AUTHORS):
        if i:
            auth_parts.append((' and ' if i == len(AUTHORS) - 1 else ', ', {}))
        auth_parts += [(name, {}), (' ' + sup, {'sup': True})]
    set_runs(pars[2], auth_parts)
    set_runs(pars[3], [(AFFILIATIONS[0], {})])
    pars[4]._p.getparent().remove(pars[4]._p)  # second affiliation placeholder
    set_runs(pars[5], [(AFFILIATIONS[1], {})])
    set_runs(pars[7], [('Abstract: ', {'bold': True}), (info['abstract'], {})])
    pars[6]._p.getparent().remove(pars[6]._p)  # standalone "Abstract" placeholder line
    set_runs(pars[8], [('Keywords: ', {'bold': True}), (info['keywords'], {})])
    line_el = pars[9]._p  # MDPI_1.9_line
    # delete template content after the line
    el = line_el.getnext()
    while el is not None and el.tag != W('w:sectPr'):
        nxt = el.getnext()
        b.body.remove(el)
        el = nxt

    # figure width = text width minus body indent
    sec = dst.sections[0]
    ind = dst.styles['MDPI_3.1_text'].paragraph_format.left_indent
    emu = (
        int(sec.page_width)
        - int(sec.left_margin)
        - int(sec.right_margin)
        - int(ind or 0)
    )
    width_cm = emu / 360000 - 0.1

    # ---- transplant body / back matter / references
    state, figs_done, first_ref = 'body', set(), True
    prev_list_num = None
    for el in src.element.body:
        tag = el.tag
        if tag == W('w:sectPr'):
            continue
        if tag == W('w:tbl'):
            b.add_table(el, src)
            continue
        text = ''.join(t.text or '' for t in el.iter(W('w:t'))).strip()
        ppr = el.find(W('w:pPr'))
        pst = ppr.find(W('w:pStyle')) if ppr is not None else None
        sname = src_style_names.get(pst.get(W('w:val')), '') if pst is not None else ''
        has_drawing = el.find('.//' + W('w:drawing')) is not None
        if not text and not has_drawing:
            continue
        numpr = ppr.find(W('w:numPr')) if ppr is not None else None
        restart = False
        if sname.startswith('Heading'):
            if text == 'Back Matter':
                state = 'back'
                continue
            if text == 'References':
                state = 'refs'
            level = int(
                sname.split()[-1]
            )  # main.md uses '##' for level 1, so pandoc levels are shifted by one
            style = {
                2: 'MDPI_2.1_heading1',
                3: 'MDPI_2.2_heading2',
                4: 'MDPI_2.3_heading3',
            }.get(level, 'MDPI_2.1_heading1')
        elif state == 'refs':
            style = 'MDPI_8.1_references'
            restart, first_ref = first_ref, False
        elif state == 'back':
            style = 'MDPI_6.2_back_matter'
        elif re.match(r'^Table \d+\.', text):
            style = 'MDPI_4.1_table_caption'
        elif numpr is not None:
            nid = numpr.find(W('w:numId')).get(W('w:val'))
            if numfmt.get(nid) == 'bullet':
                style = 'MDPI_3.8_bullet'
            else:
                style = 'MDPI_3.7_itemize'
                restart = nid != prev_list_num
                prev_list_num = nid
        else:
            style = 'MDPI_3.1_text'
        new = copy.deepcopy(el)
        b.clean_par(new, style)
        if restart:
            b.restart_num(new, style)
        b.append(new)
        if state == 'body':
            for n, (path, cap) in sorted(info['figs'].items()):
                if n not in figs_done and re.search(rf'Figure {n}\b', text):
                    b.add_figure(n, path, cap, width_cm)
                    figs_done.add(n)
    missing = set(info['figs']) - figs_done
    if missing:
        raise SystemExit(f'figures never cited: {missing}')
    OUT.parent.mkdir(exist_ok=True)
    dst.save(str(OUT))
    print(f'[OK] {OUT}')


if __name__ == '__main__':
    main()
